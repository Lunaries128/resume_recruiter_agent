import re
from functools import lru_cache
from io import BytesIO
from pathlib import Path


def clean(text):
    return "\n".join(" ".join(line.split()) for line in text.splitlines() if line.strip())


@lru_cache(maxsize=1)
def ocr_engine():
    from rapidocr_onnxruntime import RapidOCR
    return RapidOCR()


def image_text(data):
    import numpy as np
    from PIL import Image
    image = Image.open(BytesIO(data)).convert('RGB')
    result, _ = ocr_engine()(np.array(image))
    return '\n'.join(str(row[1]) for row in (result or []))


def parse_file(file_path):
    from guardrails import prepare_resume_text
    path = Path(file_path)
    metadata = {"parse_method": "direct", "ocr_page_count": 0, "failed_pages": []}
    parts = []
    if path.suffix.lower() == '.pdf':
        import fitz
        with fitz.open(path) as document:
            if len(document) > 20:
                raise ValueError('文件超过20页。')
            metadata['page_count'] = len(document)
            for index, page in enumerate(document):
                text = page.get_text('text', sort=True)
                if len(re.sub(r'\s', '', text)) < 40:
                    text = image_text(page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False).tobytes('png'))
                    metadata['ocr_page_count'] += 1
                if len(text.strip()) < 10:
                    metadata['failed_pages'].append(index+1)
                parts.append(text)
        if metadata['ocr_page_count']:
            metadata['parse_method'] = 'OCR' if metadata['ocr_page_count'] == metadata['page_count'] else '混合解析'
    elif path.suffix.lower() == '.docx':
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph
        document = Document(path)
        def read_blocks(parent):
            element = parent.element.body if hasattr(parent.element, 'body') else parent._tc
            for node in element.iterchildren():
                if node.tag.endswith('}p'):
                    parts.append(Paragraph(node, parent).text)
                elif node.tag.endswith('}tbl'):
                    table = Table(node, parent)
                    seen = set()
                    for row in table.rows:
                        for cell in row.cells:
                            if cell._tc not in seen:
                                seen.add(cell._tc)
                                read_blocks(cell)
        read_blocks(document)
        for section in document.sections:
            parts.extend(p.text for p in section.header.paragraphs)
            parts.extend(p.text for p in section.footer.paragraphs)
        metadata['parse_method'] = 'DOCX'
    elif path.suffix.lower() == '.txt':
        from charset_normalizer import from_bytes
        raw = path.read_bytes()
        try:
            text = raw.decode('utf-8-sig')
        except UnicodeDecodeError:
            match = from_bytes(raw).best()
            if match is None:
                raise ValueError('无法识别文本编码。')
            text = str(match)
        parts.append(text)
        metadata['parse_method'] = 'TXT'
    else:
        raise ValueError('不支持的文件格式。')
    raw_text = clean('\n'.join(parts))
    if len(raw_text) < 40:
        raise ValueError('未识别到足够文字，请检查原文件或转换为清晰PDF。')
    email = re.search(r'[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}', raw_text)
    phone = re.search(r'(?<!\d)1[3-9]\d{9}(?!\d)', raw_text)
    name = re.search(r'姓\s*名\s*[:：]?\s*([\u4e00-\u9fa5]{2,4})(?=\s|$)', raw_text)
    contact = {'masked_name': (name.group(1)[0]+'*') if name else '姓名未识别',
               'masked_phone': (phone.group()[:3]+'****'+phone.group()[-4:]) if phone else '',
               'email': email.group() if email else ''}
    text = prepare_resume_text(raw_text)
    text = re.sub(r'性格开朗|吃苦耐劳|学习能力强|积极参与|提升自我', '', text)
    metadata['extracted_length'] = len(raw_text)
    # 启发式质量指标，不是有标准答案验证过的OCR准确率。
    confidence = max(0, min(95, 40 + min(len(text)/30, 40)
                           + sum(word in text for word in ('教育','项目','技能','工作','学校'))*3
                           - len(metadata['failed_pages'])*15))
    return {'cleaned_text': text, 'contact': contact, 'metadata': metadata,
            'parse_confidence': round(confidence, 1)}
