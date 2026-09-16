from io import BytesIO
from pathlib import Path

import pandas as pd
import streamlit as st

from ui_common import (
    api,
    request,
    ep,
    ask_delete,
    delete_confirmation,
)


def original_page(sid, uid):
    if st.button("返回会话"):
        st.session_state.detail = None

        st.rerun()

    records = api(
        "GET",
        ep(sid, "/uploads"),
    )["uploads"]

    record = next(
        (
            item
            for item in records
            if item["id"] == uid
        ),
        None,
    )

    if not record:
        st.info(
            "上传记录已删除。"
        )
        return

    st.subheader(
        record["filename"]
    )

    response = request(
        "GET",
        ep(
            sid,
            f"/uploads/{uid}/original",
        ),
    )

    content = response.content

    st.download_button(
        "下载原文件",
        content,
        file_name=record["filename"],
    )

    suffix = Path(
        record["filename"]
    ).suffix.lower()

    if suffix == ".pdf":
        import fitz

        with fitz.open(
            stream=content,
            filetype="pdf",
        ) as document:
            if not len(document):
                st.info(
                    "PDF没有页面。"
                )
                return

            page = st.number_input(
                "页码",
                min_value=1,
                max_value=len(document),
                value=1,
                key="page_" + uid,
            )

            pixmap = document[
                int(page) - 1
            ].get_pixmap(
                matrix=fitz.Matrix(
                    1.3,
                    1.3,
                )
            )

            st.image(
                pixmap.tobytes("png"),
                use_container_width=True,
            )

    elif suffix == ".docx":
        from docx import Document

        document = Document(
            BytesIO(content)
        )

        st.caption(
            "文本预览；完整排版请下载原文件。"
        )

        for paragraph in document.paragraphs:
            if paragraph.text:
                st.text(
                    paragraph.text
                )

        for table in document.tables:
            rows = [
                [
                    cell.text
                    for cell in row.cells
                ]
                for row in table.rows
            ]

            st.dataframe(
                pd.DataFrame(rows),
                hide_index=True,
            )

    elif suffix == ".txt":
        from charset_normalizer import (
            from_bytes,
        )

        text = from_bytes(content).best()

        st.text(
            str(text)
            if text
            else "编码无法识别，请下载查看。"
        )


def uploads_page(sid):
    version_key = (
        "upload_version_" + sid
    )

    version = st.session_state.setdefault(
        version_key,
        0,
    )

    files = st.file_uploader(
        "上传简历（PDF、DOCX、TXT，每份不超过10MB）",
        type=[
            "pdf",
            "docx",
            "txt",
        ],
        accept_multiple_files=True,
        key=f"files_{sid}_{version}",
    )

    if st.button(
        "解析并入库",
        disabled=not files,
        type="primary",
    ):
        errors = []

        progress = st.progress(0)

        for index, file in enumerate(files):
            try:
                with st.spinner(
                    "正在处理：" + file.name
                ):
                    result = api(
                        "POST",
                        ep(sid, "/uploads"),
                        files={
                            "file": (
                                file.name,
                                file.getvalue(),
                                file.type,
                            )
                        },
                        timeout=75,
                    )

                if result["status"] != "成功入库":
                    errors.append(
                        file.name
                        + "："
                        + (
                            result.get("error")
                            or result["status"]
                        )
                    )

            except RuntimeError as exc:
                errors.append(
                    file.name
                    + "："
                    + str(exc)
                )

            progress.progress(
                (index + 1) / len(files)
            )

        st.session_state[version_key] += 1

        st.session_state.notice = (
            "\n".join(errors)
            or "本批次处理完成。"
        )

        st.rerun()

    records = api(
        "GET",
        ep(sid, "/uploads"),
    )["uploads"]

    if not records:
        st.info(
            "尚未上传简历，上传后可查看"
            "入库状态及原文件。"
        )
        return

    st.caption(
        "失败文件保留在上传记录中，可直接重试；"
        "置信度为文本质量估计，不是已验证准确率。"
    )

    selected = []

    # 最后一列增加宽度，避免按钮文字被挤压。
    widths = [
        0.55,
        2.6,
        1.3,
        0.9,
        2.3,
        1.25,
        1.8,
    ]

    titles = [
        "选择",
        "上传文件名",
        "入库状态",
        "置信度",
        "提示",
        "重试",
        "查看详情",
    ]

    for column, title in zip(
        st.columns(widths),
        titles,
    ):
        column.markdown(
            "**" + title + "**"
        )

    st.divider()

    for item in records:
        columns = st.columns(
            widths,
            vertical_alignment="center",
        )

        if columns[0].checkbox(
            "选择记录",
            key="select_upload_" + item["id"],
            label_visibility="collapsed",
        ):
            selected.append(
                item["id"]
            )

        columns[1].write(
            item["filename"]
        )

        columns[2].write(
            item["status"]
        )

        confidence = float(
            item.get("confidence")
            or 0
        )

        columns[3].write(
            f"{confidence:.1f}%"
        )

        columns[4].write(
            item.get("error")
            or "—"
        )

        if item["status"] not in (
            "成功入库",
            "解析中",
        ):
            if columns[5].button(
                "重新解析",
                key="retry_" + item["id"],
                use_container_width=True,
            ):
                with st.spinner(
                    "重新解析中"
                ):
                    api(
                        "POST",
                        ep(
                            sid,
                            f"/uploads/{item['id']}/retry",
                        ),
                        timeout=75,
                    )

                st.rerun()

        if columns[6].button(
            "查看原文件",
            key="original_" + item["id"],
            use_container_width=True,
        ):
            st.session_state.detail = (
                "original",
                item["id"],
            )

            st.rerun()

        st.divider()

    if st.button(
        "删除所选上传记录",
        disabled=not selected,
    ):
        ask_delete(
            sid,
            "uploads",
            selected,
            "uploads",
        )

    # 删除确认只显示在本页面，
    # 并紧跟在删除按钮下方。
    delete_confirmation("uploads")