import base64
import inspect
import re
from datetime import datetime, timedelta, timezone
from io import BytesIO
from pathlib import Path

import pandas as pd
import plotly.graph_objects as go
import requests
import streamlit as st
from PIL import Image

from config import API_URL


ROOT = Path(__file__).resolve().parent

LOGO = next(
    (
        path
        for path in (
            ROOT / 'assets/logo.svg',
            ROOT / 'asserts/logo.svg',
        )
        if path.exists()
    ),
    None,
)


def favicon():
    if LOGO:
        match = re.search(
            r'data:image/png;base64,([^\"\x27]+)',
            LOGO.read_text(encoding='utf-8'),
        )
        if match:
            return Image.open(
                BytesIO(base64.b64decode(match.group(1)))
            )
        return str(LOGO)

    return '💼'


st.set_page_config(
    page_title='智能简历筛选与招聘助理',
    page_icon=favicon(),
    layout='wide',
)

st.markdown(
    '''
    <style>
    [data-testid="stMetric"] {
        background:white;
        border:1px solid #D9E4F2;
        border-radius:10px;
        padding:12px;
    }

    button,
    button p,
    button span,
    button div[data-testid="stMarkdownContainer"] {
        white-space:nowrap!important;
        overflow:visible!important;
        text-overflow:clip!important;
        -webkit-line-clamp:unset!important;
        overflow-wrap:normal!important;
        word-break:keep-all!important;
    }

    button {
        min-width:max-content!important;
        flex-shrink:0!important;
        height:auto!important;
        min-height:2.5rem;
    }

    [data-testid="stButton"],
    [data-testid="stDownloadButton"],
    [data-testid="stFormSubmitButton"] {
        overflow-x:auto;
        max-width:100%;
    }
    </style>
    ''',
    unsafe_allow_html=True,
)

if (
    'on_change' not in inspect.signature(st.tabs).parameters
    or 'wrap' not in inspect.signature(st.container).parameters
):
    st.error(
        '请先升级：python -m pip install "streamlit>=1.63.0"，'
        '然后重启前端。'
    )
    st.stop()

if LOGO:
    tab_icon = favicon()

    if isinstance(tab_icon, Image.Image):
        tab_icon = tab_icon.copy()
        tab_icon.thumbnail((32, 32))
        buffer = BytesIO()
        tab_icon.save(buffer, format='PNG')
        icon_uri = (
            'data:image/png;base64,'
            + base64.b64encode(
                buffer.getvalue()
            ).decode('ascii')
        )
    else:
        icon_uri = (
            'data:image/svg+xml;base64,'
            + base64.b64encode(
                LOGO.read_bytes()
            ).decode('ascii')
        )

    st.markdown(
        f'''
        <style>
        button[data-baseweb="tab"]::before {{
            content:"";
            display:inline-block;
            width:18px;
            height:18px;
            margin-right:7px;
            background:url("{icon_uri}") center/contain no-repeat;
        }}
        </style>
        ''',
        unsafe_allow_html=True,
    )


def request(method, path, **kwargs):
    timeout = kwargs.pop('timeout', 20)

    try:
        response = requests.request(
            method,
            API_URL.rstrip('/') + path,
            timeout=timeout,
            **kwargs,
        )
    except requests.Timeout as exc:
        raise RuntimeError(
            '请求超时。请刷新查看服务器已保存的状态，再决定是否重试。'
        ) from exc
    except requests.RequestException as exc:
        raise RuntimeError(
            '无法连接后端，请检查后端是否已启动。'
        ) from exc

    if not response.ok:
        try:
            detail = response.json().get(
                'detail',
                response.text,
            )
        except ValueError:
            detail = response.text

        raise RuntimeError(str(detail))

    return response


def api(method, path, **kwargs):
    return request(method, path, **kwargs).json()


def endpoint(sid, tail=''):
    return f'/sessions/{sid}{tail}'


def open_session(sid):
    st.session_state.active_session = sid
    st.session_state.page = 'workspace'
    st.session_state.detail = None
    st.rerun()


def confirm_request(
    sid,
    kind,
    targets=None,
    location='sidebar',
):
    try:
        result = api(
            'POST',
            endpoint(sid, '/delete/request'),
            json={
                'kind': kind,
                'targets': targets or [],
            },
        )

        st.session_state.pending_delete = {
            'sid': sid,
            'location': location,
            **result,
        }

        st.rerun()

    except RuntimeError as exc:
        st.error(str(exc))


def delete_confirmation(location):
    pending = st.session_state.get('pending_delete')

    if not pending or pending.get('location') != location:
        return

    if (
        location in ('uploads', 'candidates')
        and pending['sid'] != st.session_state.active_session
    ):
        return

    with st.container(border=True):
        label = (
            '整个会话及其全部原文件、候选人、评分和聊天记录'
            if pending['kind'] == 'session'
            else (
                f"所选的 {pending['count']} 条上传记录"
                '及其原文件和候选人评分'
            )
        )

        st.warning(
            '确认删除' + label + '？此操作不可恢复。'
        )

        accepted = st.checkbox(
            '我确认删除',
            key='confirm_' + pending['token'],
        )

        bar = st.container(
            horizontal=True,
            wrap=False,
            horizontal_alignment='left',
            gap='small',
        )

        if bar.button(
            '确认删除',
            disabled=not accepted,
            key='delete_yes',
        ):
            try:
                result = api(
                    'POST',
                    endpoint(
                        pending['sid'],
                        '/delete/confirm',
                    ),
                    json={
                        'token': pending['token'],
                        'confirmed': True,
                    },
                )

                if (
                    result['kind'] == 'session'
                    and st.session_state.active_session
                    == pending['sid']
                ):
                    st.session_state.active_session = None
                    st.session_state.page = 'workspace'

                if result['kind'] == 'session':
                    sid = pending['sid']

                    for key in list(st.session_state):
                        if key in (
                            f'jd_{sid}',
                            f'draft_jd_{sid}',
                            f'match_{sid}',
                            f'upload_version_{sid}',
                            f'module_{sid}',
                            f'active_module_{sid}',
                        ):
                            st.session_state.pop(key, None)

                st.session_state.pending_delete = None
                st.session_state.detail = None
                st.session_state.notice = (
                    '删除完成。'
                    + ' '.join(result.get('warnings', []))
                )
                st.rerun()

            except RuntimeError as exc:
                st.error(str(exc))

        if bar.button('取消', key='delete_no'):
            st.session_state.pending_delete = None
            st.rerun()


def is_today(value):
    local = timezone(timedelta(hours=8))

    return (
        datetime.fromisoformat(value).astimezone(local).date()
        == datetime.now(local).date()
    )


def preference_panel():
    hr = st.session_state.hr_id

    with st.expander('长期偏好管理'):
        st.caption(
            '保存后跨会话生效；仅将岗位相关偏好作为优先项。'
        )

        with st.form('new_pref'):
            content = st.text_area('偏好内容')
            scope = st.text_input(
                '适用类别',
                value='通用',
            )

            if st.form_submit_button('保存偏好'):
                api(
                    'POST',
                    '/preferences',
                    json={
                        'hr_id': hr,
                        'content': content,
                        'scope': scope,
                    },
                )
                st.rerun()

        saved = api(
            'GET',
            '/preferences',
            params={'hr_id': hr},
        )['preferences']

        for item in saved:
            with st.form('pref_' + item['id']):
                content = st.text_area(
                    '偏好',
                    value=item['content'],
                )
                scope = st.text_input(
                    '类别',
                    value=item['scope'],
                )
                enabled = st.checkbox(
                    '启用',
                    value=bool(item['enabled']),
                )

                if st.form_submit_button('更新偏好'):
                    api(
                        'PUT',
                        '/preferences/' + item['id'],
                        json={
                            'hr_id': hr,
                            'content': content,
                            'scope': scope,
                            'enabled': enabled,
                        },
                    )
                    st.rerun()

                sure = st.checkbox('确认删除此偏好')

                if st.form_submit_button('删除偏好'):
                    if not sure:
                        st.warning('请先勾选确认删除此偏好。')
                    else:
                        api(
                            'DELETE',
                            '/preferences/' + item['id'],
                            params={
                                'hr_id': hr,
                                'confirmed': True,
                            },
                        )
                        st.rerun()


def sidebar(sessions):
    with st.sidebar:
        logo_col, text_col = st.columns(
            [1, 4],
            vertical_alignment='center',
        )

        if LOGO:
            logo_col.image(str(LOGO), width=40)

        text_col.markdown('### 智能小助理')
        st.text_input('HR编号', key='hr_id')
        preference_panel()

        if st.button(
            '＋ 新建招聘会话',
            use_container_width=True,
        ):
            try:
                open_session(
                    api('POST', '/sessions')['id']
                )
            except RuntimeError as exc:
                st.error(str(exc))

        if st.button(
            '历史文件夹',
            use_container_width=True,
        ):
            st.session_state.page = 'history'
            st.session_state.detail = None
            st.rerun()

        st.divider()

        groups = (
            (
                '置顶会话',
                [
                    item
                    for item in sessions
                    if item['pinned']
                ],
            ),
            (
                '今天的会话',
                [
                    item
                    for item in sessions
                    if not item['pinned']
                    and is_today(item['created_at'])
                ],
            ),
        )

        for title, items in groups:
            st.caption(title)

            if not items:
                st.caption('暂无')

            for item in items:
                if st.button(
                    item['title'],
                    key='open_' + item['id'],
                    use_container_width=True,
                ):
                    open_session(item['id'])

        sid = st.session_state.active_session
        current = next(
            (
                item
                for item in sessions
                if item['id'] == sid
            ),
            None,
        )

        if current:
            st.divider()

            with st.expander('当前会话设置'):
                st.caption(current['title'])

                with st.form('rename_' + sid):
                    name = st.text_input(
                        '会话名称',
                        value=current['title'],
                        max_chars=80,
                    )

                    if st.form_submit_button('保存名称'):
                        try:
                            api(
                                'PATCH',
                                endpoint(sid),
                                json={'title': name},
                            )
                            st.rerun()
                        except RuntimeError as exc:
                            st.error(str(exc))

                if st.button(
                    '取消置顶'
                    if current['pinned']
                    else '置顶当前会话'
                ):
                    api(
                        'PATCH',
                        endpoint(sid),
                        json={
                            'pinned': not bool(
                                current['pinned']
                            ),
                        },
                    )
                    st.rerun()

                if st.button('删除当前会话'):
                    confirm_request(
                        sid,
                        'session',
                        location='sidebar',
                    )

            delete_confirmation('sidebar')

            if st.button(
                '清空当前对话',
                use_container_width=True,
            ):
                api(
                    'POST',
                    endpoint(sid, '/clear'),
                )
                st.rerun()


def history_page(sessions):
    st.title('历史文件夹')
    st.caption(
        '所有会话均保存在本机。打开历史会话可继续查看'
        '原文件、候选人和匹配结果。'
    )

    query = st.text_input(
        '搜索会话名称'
    ).strip().lower()

    items = [
        item
        for item in sessions
        if query in item['title'].lower()
    ]

    if not items:
        st.info('暂无历史会话，请在侧边栏新建会话。')

    for item in items:
        with st.container(border=True):
            info, actions = st.columns([7, 3])

            bar = actions.container(
                horizontal=True,
                wrap=False,
                horizontal_alignment='right',
                gap='small',
            )

            info.markdown(
                ('📌 ' if item['pinned'] else '')
                + item['title']
            )

            created = datetime.fromisoformat(
                item['created_at']
            ).astimezone(
                timezone(timedelta(hours=8))
            )

            info.caption(
                f'创建于 {created:%Y-%m-%d %H:%M}'
                f" · {item['upload_count']} 份上传"
                f" · {item['candidate_count']} 名候选人"
            )

            if bar.button(
                '打开',
                key='history_open_' + item['id'],
            ):
                open_session(item['id'])

            if bar.button(
                '取消置顶' if item['pinned'] else '置顶',
                key='pin_' + item['id'],
            ):
                api(
                    'PATCH',
                    endpoint(item['id']),
                    json={
                        'pinned': not bool(item['pinned']),
                    },
                )
                st.rerun()

            if bar.button(
                '删除',
                key='del_' + item['id'],
            ):
                confirm_request(
                    item['id'],
                    'session',
                    location='history',
                )

    delete_confirmation('history')


def original_page(sid, uid):
    if st.button('← 返回会话'):
        st.session_state.detail = None
        st.rerun()

    records = api(
        'GET',
        endpoint(sid, '/uploads'),
    )['uploads']

    record = next(
        (
            item
            for item in records
            if item['id'] == uid
        ),
        None,
    )

    if not record:
        st.info('该上传记录已被删除。')
        return

    st.subheader(record['filename'])
    st.caption('状态：' + record['status'])

    if record['error']:
        st.warning(record['error'])

    response = request(
        'GET',
        endpoint(sid, f'/uploads/{uid}/original'),
    )

    content = response.content

    st.download_button(
        '下载原文件',
        content,
        file_name=record['filename'],
        mime=response.headers.get(
            'content-type',
            'application/octet-stream',
        ),
    )

    suffix = Path(record['filename']).suffix.lower()

    if suffix == '.pdf':
        import fitz

        with fitz.open(
            stream=content,
            filetype='pdf',
        ) as document:
            if not len(document):
                st.info('PDF没有页面。')
                return

            page = st.number_input(
                '页码',
                1,
                len(document),
                1,
                key='pdf_page_' + uid,
            )

            pix = document[int(page) - 1].get_pixmap(
                matrix=fitz.Matrix(1.4, 1.4),
                alpha=False,
            )

            st.image(
                pix.tobytes('png'),
                use_container_width=True,
            )

    elif suffix == '.txt':
        from charset_normalizer import from_bytes

        decoded = from_bytes(content).best()

        st.text(
            str(decoded)
            if decoded
            else '编码无法识别，请下载原文件查看。'
        )

    elif suffix == '.docx':
        from docx import Document

        document = Document(BytesIO(content))

        st.caption(
            '下方为原DOCX的文本和表格预览；'
            '要查看完整排版、图片和文本框，请下载原文件。'
        )

        for paragraph in document.paragraphs:
            if paragraph.text:
                st.text(paragraph.text)

        for table in document.tables:
            st.dataframe(
                pd.DataFrame([
                    [
                        cell.text
                        for cell in row.cells
                    ]
                    for row in table.rows
                ]),
                hide_index=True,
                use_container_width=True,
            )

    else:
        st.info('此格式无法在线预览，可以下载原文件。')


def uploads_page(sid):
    version_key = 'upload_version_' + sid
    version = st.session_state.setdefault(
        version_key,
        0,
    )

    files = st.file_uploader(
        '上传简历（PDF、DOCX、TXT，每份不超过10MB）',
        type=['pdf', 'docx', 'txt'],
        accept_multiple_files=True,
        key=f'files_{sid}_{version}',
    )

    if st.button(
        '解析并入库',
        disabled=not files,
        type='primary',
    ):
        errors = []
        progress = st.progress(0)

        for index, file in enumerate(files):
            with st.spinner('正在处理：' + file.name):
                try:
                    api(
                        'POST',
                        endpoint(sid, '/uploads'),
                        files={
                            'file': (
                                file.name,
                                file.getvalue(),
                                file.type,
                            ),
                        },
                        timeout=75,
                    )
                except RuntimeError as exc:
                    errors.append(
                        file.name + '：' + str(exc)
                    )

            progress.progress(
                (index + 1) / len(files)
            )

        st.session_state[version_key] += 1
        st.session_state.notice = (
            '\n'.join(errors)
            if errors
            else '已提交后台解析，请点击刷新查看结果。'
        )
        st.rerun()

    if st.button(
        '刷新解析进度',
        key='refresh_' + sid,
    ):
        st.rerun()

    records = api(
        'GET',
        endpoint(sid, '/uploads'),
    )['uploads']

    if not records:
        st.info(
            '尚未上传简历。上传后会在这里显示解析状态，'
            '并可查看原文件。'
        )
        return

    st.caption(
        '解析置信度为文本质量估计，不代表已验证的准确率。'
        '解析失败或超时的原文件仍然保留。'
    )

    widths = [0.5, 2.5, 1.3, 0.8, 2, 1, 1.8]

    for col, title in zip(
        st.columns(widths),
        [
            '选择',
            '上传文件名',
            '入库状态',
            '置信度',
            '提示',
            '重试',
            '查看详情',
        ],
    ):
        col.markdown('**' + title + '**')

    selected = []

    for item in records:
        cols = st.columns(
            widths,
            vertical_alignment='center',
        )

        if cols[0].checkbox(
            '选择',
            key='upload_select_' + item['id'],
            label_visibility='collapsed',
        ):
            selected.append(item['id'])

        cols[1].write(item['filename'])
        cols[2].write(item['status'])
        cols[3].write(
            f"{item['confidence']:.1f}%"
        )
        cols[4].write(item['error'] or '—')

        if item['status'] not in (
            '成功入库',
            '解析中',
            '等待解析',
        ):
            if cols[5].button(
                '重试',
                key='retry_' + item['id'],
            ):
                with st.spinner('重新解析中…'):
                    api(
                        'POST',
                        endpoint(
                            sid,
                            f"/uploads/{item['id']}/retry",
                        ),
                        timeout=75,
                    )
                st.rerun()

        if cols[6].button(
            '查看原文件',
            key='original_' + item['id'],
        ):
            st.session_state.detail = (
                'original',
                item['id'],
            )
            st.rerun()

        st.divider()

    if st.button(
        '删除所选上传记录',
        disabled=not selected,
    ):
        confirm_request(
            sid,
            'uploads',
            selected,
            location='uploads',
        )

    delete_confirmation('uploads')


def remember_jd(sid):
    st.session_state['draft_jd_' + sid] = (
        st.session_state['jd_' + sid]
    )


def job_page(session):
    sid = session['id']

    st.session_state.setdefault(
        'draft_jd_' + sid,
        session['jd'],
    )
    st.session_state.setdefault(
        'jd_' + sid,
        st.session_state['draft_jd_' + sid],
    )

    jd = st.text_area(
        '招聘要求 JD',
        height=190,
        key='jd_' + sid,
        on_change=remember_jd,
        args=(sid,),
    )

    scope = st.text_input(
        '岗位类别（用于偏好适用范围）',
        value='通用',
        key='scope_' + sid,
    )

    use_memory = st.checkbox(
        '使用本HR在此类别及通用范围的已启用偏好',
        value=True,
        key='memory_' + sid,
    )

    if st.button(
        '整理岗位要求',
        disabled=not jd.strip(),
    ):
        with st.spinner('整理要求与历史偏好…'):
            api(
                'POST',
                endpoint(sid, '/criteria/prepare'),
                json={
                    'jd': jd,
                    'hr_id': st.session_state.hr_id,
                    'scope': scope,
                    'use_memory': use_memory,
                },
                timeout=120,
            )

        st.session_state.pop(
            'match_' + sid,
            None,
        )
        st.rerun()

    state = api(
        'GET',
        endpoint(sid, '/criteria'),
    )
    spec = state['criteria']

    if (
        spec
        and spec['jd_text'] == jd.strip()
        and spec.get('scoring_version')
        == 'dimension-relative-v1'
    ):
        types = [
            '技能',
            '经验',
            '学历',
            '项目',
            '加分项（优先条件）',
            '需确认',
        ]

        defaults = dict(
            zip(
                types[:5],
                [50.0, 25.0, 15.0, 10.0, 0.0],
            )
        )

        saved_weights = spec.get(
            'dimension_weights',
            defaults,
        )

        version = (
            sid
            + '_'
            + str(
                spec.get('revision')
                or spec.get('draft_id', 'draft')
            )
        )

        st.caption(
            '维度权重：五项合计100%。'
            '在各维度内部，再按每条岗位要求的相对权重分配。'
        )

        weights = {}

        for index, col in enumerate(st.columns(5)):
            category = types[index]

            weights[category] = col.number_input(
                category + '（%）',
                min_value=0.0,
                max_value=100.0,
                value=float(saved_weights[category]),
                step=1.0,
                format='%.2f',
                key=(
                    'dimension_'
                    + version
                    + '_'
                    + str(index)
                ),
            )

        total = sum(weights.values())
        valid_total = abs(total - 100.0) <= 0.001

        st.caption(
            f'当前维度权重合计：{total:.2f}%'
        )

        if not valid_total:
            st.warning(
                '五个维度的权重之和必须为100%，请调整后确认。'
            )

        st.caption(
            '下表保留逐条拆分，可直接修改类型、相对权重、'
            '来源原文，并在底部新增岗位要求。'
            '技能、经验、学历、项目的初始相对权重为1；'
            '加分项和需确认为0。'
            '新增行留空的相对权重在确认时按类型补默认值；'
            '手动填0会保留。'
        )

        frame = pd.DataFrame(
            spec['requirements']
        )[
            ['name', 'category', 'weight', 'quote']
        ]

        edited = st.data_editor(
            frame,
            num_rows='dynamic',
            hide_index=True,
            use_container_width=True,
            key='criteria_' + version,
            column_config={
                'category': st.column_config.SelectboxColumn(
                    '类型',
                    options=types,
                    required=True,
                ),
                'name': st.column_config.TextColumn(
                    '岗位要求',
                    required=True,
                ),
                'weight': st.column_config.NumberColumn(
                    '相对权重',
                    min_value=0.0,
                    max_value=100.0,
                    step=0.1,
                ),
                'quote': '来源原文',
            },
        )

        records = edited.fillna('').to_dict('records')

        dirty = (
            records
            != frame.fillna('').to_dict('records')
            or weights != saved_weights
        )

        if dirty:
            st.warning(
                '岗位要求或权重已修改，请重新确认后评分。'
            )

        row_totals = {
            key: 0.0
            for key in types[:5]
        }
        rows_valid = True

        for row in records:
            if not str(
                row.get('name') or ''
            ).strip():
                continue

            category = row.get('category')

            if category not in types:
                rows_valid = False
                continue

            value = row.get('weight')

            if value == '' or value is None:
                value = (
                    1.0 if category in types[:4] else 0.0
                )
            else:
                value = float(value)

            if not 0 <= value <= 100:
                rows_valid = False
            elif category in row_totals:
                row_totals[category] += value

        active = {
            key: weights[key]
            for key in types[:5]
            if row_totals[key] > 0
            and weights[key] > 0
        }

        active_total = sum(active.values())

        if active_total:
            st.caption(
                '实际参与评分的维度：'
                + '、'.join(
                    f'{key} {value / active_total * 100:.2f}%'
                    for key, value in active.items()
                )
            )
        else:
            st.warning(
                '请至少保留一个维度权重和条目相对权重'
                '均大于0的岗位要求。'
            )

        st.caption(
            '没有有效要求的维度不扣分，其余参与维度归一化。'
            '加分项只有两层权重均大于0才计分，不作为硬门槛。'
            '需确认的相对权重可编辑并保存，明确类型前始终不参与评分。'
            '修改已有条目的类型会保留你填写的相对权重，'
            '可继续手动调整。'
        )

        invalid = (
            not valid_total
            or not rows_valid
            or active_total <= 0
        )

        if st.button(
            '确认招聘标准与权重',
            disabled=invalid,
        ):
            api(
                'POST',
                endpoint(sid, '/criteria/confirm'),
                json={
                    'requirements': records,
                    'dimension_weights': weights,
                },
            )

            st.session_state.pop(
                'match_' + sid,
                None,
            )
            st.rerun()

        st.caption(
            '状态：'
            + (
                '已确认'
                if state['confirmed']
                else '待确认'
            )
        )

        count = len(
            api(
                'GET',
                endpoint(sid, '/candidates'),
            )['candidates']
        )

        if st.button(
            '开始匹配本会话候选人',
            disabled=(
                not state['confirmed']
                or dirty
                or invalid
                or not count
            ),
            type='primary',
        ):
            with st.spinner(
                '正在逐条核验并计算权重…'
            ):
                result = api(
                    'POST',
                    endpoint(sid, '/score'),
                    json={'jd': jd},
                    timeout=(
                        10,
                        120 * (2 * count + 1) + 60,
                    ),
                )

            st.session_state[
                'match_' + sid
            ] = result

        result = st.session_state.get(
            'match_' + sid
        )

        if (
            result
            and not dirty
            and result.get(
                'requirements',
                {},
            ).get('revision') == spec.get('revision')
        ):
            ok = sum(
                item['success']
                for item in result['results']
            )

            st.info(
                f'本次完成{ok}/'
                f'{len(result["results"])}人。'
            )

            for item in result['results']:
                if not item['success']:
                    st.error(
                        item['candidate_code']
                        + '：'
                        + item['error']
                    )

    else:
        st.info(
            '请先点击“整理岗位要求”，生成六种类型，再确认标准。'
            '旧版标准需要重新整理。'
        )

    st.caption(
        '仅使用JD、简历、本地库和已保存偏好；不联网搜索。'
    )
    st.divider()

    with st.expander(
        '招聘助理聊天',
        expanded=True,
    ):
        messages = api(
            'GET',
            endpoint(sid, '/messages'),
        )['messages']

        for message in messages:
            with st.chat_message(message['role']):
                st.markdown(message['content'])

                if message.get('trace'):
                    with st.expander('工具执行记录'):
                        st.json(message['trace'])

        with st.form(
            'chat_' + sid,
            clear_on_submit=True,
        ):
            message = st.text_input(
                '输入筛选、比较或追问要求'
            )
            send = st.form_submit_button('发送')

        if send and message.strip():
            if jd.strip() != session['jd'].strip():
                st.warning(
                    '请先整理并确认修改后的JD。'
                )
                return

            count = len(
                api(
                    'GET',
                    endpoint(sid, '/candidates'),
                )['candidates']
            )

            with st.spinner('助理正在处理…'):
                api(
                    'POST',
                    endpoint(sid, '/chat'),
                    json={
                        'message': message,
                        'hr_id': st.session_state.hr_id,
                    },
                    timeout=(
                        10,
                        180 * (count + 2),
                    ),
                )
            st.rerun()

        reports = api(
            'GET',
            endpoint(sid, '/reports'),
        )['reports']

        if reports:
            with st.expander('本会话生成的报告'):
                for filename in reports:
                    content = request(
                        'GET',
                        endpoint(
                            sid,
                            '/reports/' + filename,
                        ),
                    ).content

                    st.download_button(
                        '下载报告 ' + filename[:8],
                        content,
                        file_name=filename,
                        mime='text/markdown',
                        key='report_' + filename,
                    )


def education(candidate):
    levels = {
        '博士': 5,
        '硕士': 4,
        '研究生': 4,
        '本科': 3,
        '学士': 3,
        '大专': 2,
        '专科': 2,
        '高中': 1,
    }

    items = candidate.get('education', [])

    def level(item):
        return max(
            [
                value
                for key, value in levels.items()
                if key in item.get('degree', '')
            ]
            or [0]
        )

    return max(items, key=level) if items else {}


def experience(candidate):
    facts = []

    if candidate.get('work_years'):
        facts.append(
            f"{candidate['work_years']}年工作经验"
        )

    items = (
        candidate.get('experiences', [])
        + candidate.get('projects', [])
    )

    if items:
        item = items[0]

        facts += [
            item.get('role')
            or item.get('name')
            or item.get('organization')
            or ''
        ]

        facts += (
            item.get('actions', [])
            + item.get('results', [])
        )[:2]

    return (
        '；'.join(item for item in facts if item)
        or '信息缺失'
    )


def candidate_page(sid):
    candidates = api(
        'GET',
        endpoint(sid, '/candidates'),
    )['candidates']

    if not candidates:
        st.info(
            '本会话尚无成功入库的简历，'
            '请先前往“简历入库”上传。'
        )
        return

    st.caption(
        '按已获证据支持分降序，同分比较证据覆盖率；'
        '明确差距及未知项需人工核实。'
    )

    widths = [
        0.5, 0.6, 1, 2, 1.7, 0.9, 2.5, 1.5
    ]

    for col, label in zip(
        st.columns(widths),
        [
            '选择',
            '排名',
            '姓名',
            '匹配概况',
            '学校专业',
            '最高学历',
            '经历',
            '操作',
        ],
    ):
        col.markdown('**' + label + '**')

    selected = []

    for candidate in candidates:
        code = candidate['candidate_code']

        cols = st.columns(
            widths,
            vertical_alignment='center',
        )

        if cols[0].checkbox(
            '选择',
            key='candidate_select_' + code,
            label_visibility='collapsed',
        ):
            selected.append(candidate['upload_id'])

        cols[1].write(
            candidate.get('rank') or '未评分'
        )

        cols[2].write(
            candidate.get('masked_name')
            or '姓名未识别'
        )

        detail = candidate.get('score_detail') or {}

        cols[3].write(
            detail.get('recommendation')
            or (
                '旧版结果，请重新匹配'
                if detail
                else '未核验'
            )
        )

        if detail.get('summary'):
            cols[3].caption(
                f"支持分{detail['total_score']}"
                f" / 覆盖率{detail['coverage']}%"
            )
            cols[3].caption(detail['summary'])

        edu = education(candidate)

        cols[4].write(
            ' / '.join(
                item
                for item in (
                    edu.get('school'),
                    edu.get('major'),
                )
                if item
            )
            or '信息缺失'
        )

        cols[5].write(
            edu.get('degree') or '未知'
        )
        cols[6].write(experience(candidate))

        if cols[7].button(
            '查看详情',
            key='candidate_detail_' + code,
        ):
            st.session_state.detail = (
                'candidate',
                code,
            )
            st.rerun()

        st.divider()

    if st.button(
        '删除所选候选人及原文件',
        disabled=not selected,
    ):
        confirm_request(
            sid,
            'uploads',
            selected,
            location='candidates',
        )

    delete_confirmation('candidates')
    st.divider()

    st.caption(
        '下载报告：勾选时生成所选候选人报告；'
        '未勾选时生成本会话全部候选人报告。'
        '请先完成当前JD匹配。'
    )

    if st.button(
        '下载报告',
        key='generate_pdf_' + sid,
    ):
        codes = [
            item['candidate_code']
            for item in candidates
            if not selected
            or item['upload_id'] in selected
        ]

        with st.spinner('正在生成报告…'):
            result = api(
                'POST',
                endpoint(sid, '/screening-reports'),
                json={'candidate_codes': codes},
                timeout=120,
            )

        st.session_state.detail = (
            'report',
            result['filename'],
        )
        st.rerun()

    saved = api(
        'GET',
        endpoint(sid, '/screening-reports'),
    )['reports']

    if saved:
        with st.expander('已归档PDF报告'):
            for filename in saved:
                if st.button(
                    '预览报告 ' + filename[:8],
                    key='pdf_' + filename,
                ):
                    st.session_state.detail = (
                        'report',
                        filename,
                    )
                    st.rerun()


def report_page(sid, filename):
    if st.button('← 返回候选人信息'):
        st.session_state[
            'active_module_' + sid
        ] = '候选人信息'
        st.session_state.detail = None
        st.rerun()

    report = api(
        'GET',
        endpoint(
            sid,
            '/screening-reports/' + filename,
        ),
    )

    content = request(
        'GET',
        endpoint(
            sid,
            '/reports/' + filename,
        ),
    ).content

    st.title(report['title'])

    st.download_button(
        '下载PDF报告',
        content,
        file_name='招聘筛选报告_' + filename,
        mime='application/pdf',
        key='download_' + filename,
    )

    st.caption(
        '生成时间：' + report['created_at']
    )
    st.info(report['note'])

    if len(report.get('overview', [])) > 1:
        st.subheader('候选人汇总')
        st.dataframe(
            pd.DataFrame(report['overview']),
            hide_index=True,
            use_container_width=True,
        )

    for candidate in report['candidates']:
        st.header(
            candidate['name']
            + ' / '
            + candidate['code']
        )

        if candidate.get('job_title'):
            st.caption(
                '岗位：' + candidate['job_title']
            )

        for section in candidate['sections']:
            st.subheader(section['title'])

            for index, item in enumerate(
                section['items'],
                1,
            ):
                st.text(f'{index}. {item}')

            if (
                section['title'] == '初筛结论'
                and candidate.get('checks')
            ):
                st.subheader('岗位要求核验')
                st.dataframe(
                    pd.DataFrame([
                        {
                            '要求': item['requirement'],
                            '类型': item['category'],
                            '判断': item['status'],
                            '依据': item['evidence'],
                        }
                        for item in candidate['checks']
                    ]),
                    hide_index=True,
                    use_container_width=True,
                )

        st.divider()

    with st.expander('报告版本与岗位原文'):
        st.caption(
            '报告编号：'
            + report['id']
            + ' / 版本：'
            + report['version']
        )
        st.caption(
            '会话编号：' + report['session_id']
        )
        st.caption(
            '岗位版本SHA-256：' + report['jd_hash']
        )
        st.text(report['jd'])


def review_panel(detail):
    st.subheader(detail['recommendation'])
    st.write(detail['summary'])

    st.caption(
        '以上为人工初筛建议；简历陈述未经独立核实。'
    )

    rows = [
        {
            '岗位要求': item['name'],
            '类型': item['category'],
            '判断': item['status'],
            '依据 / 待核实': item['reason'],
        }
        for item in detail['reviews']
    ]

    st.dataframe(
        pd.DataFrame(rows),
        hide_index=True,
        use_container_width=True,
    )

    with st.expander('证据原文与针对性面试问题'):
        for item in detail['reviews']:
            st.markdown(
                '**'
                + item['name']
                + ' / '
                + item['status']
                + '**'
            )

            for ref in item['references']:
                st.text(
                    ref['source']
                    + '：'
                    + ref['quote']
                )

            st.write(
                '建议提问：' + item['question']
            )

    with st.expander('辅助分数说明'):
        low, high = detail['score_range']

        st.write(
            f'已获证据支持分：{low:.1f}/100；'
            f'证据覆盖率：{detail["coverage"]:.1f}%。'
        )
        st.write(
            f'若未知核心条件均满足，区间上界为{high:.1f}；'
            '这不是胜任概率。'
        )
        st.caption(detail['scoring_rule'])


def candidate_detail(sid, code):
    if st.button('← 返回候选人信息'):
        st.session_state[
            'active_module_' + sid
        ] = '候选人信息'
        st.session_state.detail = None
        st.rerun()

    candidates = api(
        'GET',
        endpoint(sid, '/candidates'),
    )['candidates']

    candidate = next(
        (
            item
            for item in candidates
            if item['candidate_code'] == code
        ),
        None,
    )

    if not candidate:
        st.info('当前会话中不存在该候选人。')
        return

    left, right = st.columns(
        [6, 4],
        gap='large',
    )

    with left:
        st.subheader(
            candidate.get('masked_name')
            or '姓名未识别'
        )
        st.write(
            '电话：'
            + (
                candidate.get('masked_phone')
                or '未提供'
            )
        )
        st.write(
            '邮箱：'
            + (
                candidate.get('email')
                or '未提供'
            )
        )

        if st.button('查看原文件'):
            st.session_state.detail = (
                'original',
                candidate['upload_id'],
            )
            st.rerun()

        for title, key in [
            ('教育经历', 'education'),
            ('工作与实习', 'experiences'),
            ('项目经历', 'projects'),
            ('竞赛获奖', 'awards'),
        ]:
            st.subheader(title)
            items = candidate.get(key, [])

            if not items:
                st.caption('未提供')

            for item in items:
                header = ' · '.join(
                    str(item.get(field, ''))
                    for field in (
                        'school',
                        'major',
                        'degree',
                        'organization',
                        'role',
                        'name',
                        'level',
                    )
                    if item.get(field)
                )
                st.write(header)

                dates = ' — '.join(
                    str(item.get(field, ''))
                    for field in (
                        'start_date',
                        'end_date',
                        'date',
                    )
                    if item.get(field)
                )

                if dates:
                    st.caption(dates)

                for fact in (
                    item.get('actions', [])
                    + item.get('results', [])
                ):
                    st.write('• ' + fact)

                if item.get('technologies'):
                    st.caption(
                        '技术：'
                        + '、'.join(item['technologies'])
                    )

        st.subheader('技能与证书')
        st.write(
            '、'.join(
                candidate.get('skills', [])
                + candidate.get('certificates', [])
            )
            or '未提供'
        )

    with right:
        detail = candidate.get(
            'score_detail',
            {},
        )

        if not detail:
            st.info('当前JD尚未评分。')
            return

        if detail.get('matching_version') == 'evidence-v2':
            review_panel(detail)
        else:
            st.warning(
                '这是旧版评分，请在招聘要求页重新匹配。'
            )
            st.metric(
                '旧版匹配分',
                detail['total_score'],
            )

        with st.expander('查看辅助图表'):
            dimensions = detail.get(
                'dimensions',
                [],
            )

            if (
                detail.get('matching_version')
                == 'evidence-v2'
            ):
                known = {
                    item['id']
                    for item in detail['reviews']
                    if item['status'] != '信息不足'
                }

                dimensions = [
                    item
                    for item in dimensions
                    if item['key'] in known
                ]

                st.caption(
                    '图表仅显示证据足以作出判断的核心条件；'
                    '未知条件不画成零分。'
                )

            kind = st.radio(
                '图表类型',
                ['雷达图', '柱状图'],
                horizontal=True,
            )

            if dimensions:
                names = [
                    item['name']
                    for item in dimensions
                ]
                values = [
                    item['score']
                    for item in dimensions
                ]

                if kind == '雷达图':
                    fig = go.Figure(
                        go.Scatterpolar(
                            r=values + [values[0]],
                            theta=names + [names[0]],
                            fill='toself',
                            line_color='#82A5D9',
                        )
                    )
                    fig.update_layout(
                        polar={
                            'radialaxis': {
                                'range': [0, 100],
                            },
                        },
                    )
                else:
                    fig = go.Figure(
                        go.Bar(
                            x=names,
                            y=values,
                            marker_color='#82A5D9',
                        )
                    )
                    fig.update_yaxes(range=[0, 100])

                fig.update_layout(
                    showlegend=False,
                    margin={
                        'l': 35,
                        'r': 35,
                        't': 20,
                        'b': 20,
                    },
                )
                st.plotly_chart(
                    fig,
                    use_container_width=True,
                )

        with st.expander('查看完整评分审计记录'):
            for line in detail.get('audit_log', []):
                st.write(line)


def remember_module(sid):
    st.session_state[
        'active_module_' + sid
    ] = st.session_state[
        'module_' + sid
    ]


def main():
    for key, value in {
        'active_session': None,
        'page': 'workspace',
        'detail': None,
        'hr_id': 'default_hr',
    }.items():
        st.session_state.setdefault(key, value)

    sessions = api(
        'GET',
        '/sessions',
    )['sessions']

    if st.session_state.active_session not in [
        item['id']
        for item in sessions
    ]:
        st.session_state.active_session = None

    if (
        st.session_state.page == 'workspace'
        and not st.session_state.active_session
    ):
        st.session_state.active_session = api(
            'POST',
            '/sessions',
        )['id']

        sessions = api(
            'GET',
            '/sessions',
        )['sessions']

    sidebar(sessions)

    notice = st.session_state.pop(
        'notice',
        None,
    )

    if notice:
        st.info(notice)

    if (
        st.session_state.page == 'history'
        or not st.session_state.active_session
    ):
        history_page(sessions)
        return

    sid = st.session_state.active_session
    session = api('GET', endpoint(sid))
    detail = st.session_state.get('detail')

    if detail:
        pages = {
            'original': original_page,
            'candidate': candidate_detail,
            'report': report_page,
        }
        pages[detail[0]](sid, detail[1])
        return

    st.title('智能简历筛选与招聘助理')
    st.caption(
        '当前会话：' + session['title']
    )

    st.session_state.setdefault(
        'active_module_' + sid,
        '简历入库',
    )
    st.session_state.setdefault(
        'module_' + sid,
        st.session_state['active_module_' + sid],
    )

    tabs = st.tabs(
        [
            '简历入库',
            '招聘要求',
            '候选人信息',
        ],
        key='module_' + sid,
        on_change=remember_module,
        args=(sid,),
    )

    with tabs[0]:
        uploads_page(sid)

    with tabs[1]:
        job_page(session)

    with tabs[2]:
        candidate_page(sid)


try:
    main()
except RuntimeError as exc:
    st.error(str(exc))
except Exception as exc:
    st.error(
        f'页面处理失败：{type(exc).__name__}: {exc}'
    )